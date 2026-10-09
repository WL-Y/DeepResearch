"""Local, in-memory report conversion. Never fetch resources from Markdown."""

from functools import lru_cache
from html import escape
from io import BytesIO
from pathlib import Path
import os
import re
from urllib.parse import urlsplit

from markdown_it import MarkdownIt


MIME_TYPES = {
    "md": "text/markdown; charset=utf-8",
    "txt": "text/plain; charset=utf-8",
    "html": "text/html; charset=utf-8",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}


def safe_filename(topic: str, extension: str) -> str:
    name = re.sub(r'[\\/:*?"<>|\x00-\x1f\x7f]', '-', topic).strip(' .')[:80].rstrip(' .')
    if not name or re.match(r'^(CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(?:\.|$)', name, re.I):
        name = 'research-report'
    return f'{name}.{extension}'


def safe_url(url: str) -> bool:
    try:
        parsed = urlsplit(url)
        return parsed.scheme.lower() in {'http', 'https', 'mailto'} and not re.search(r'[\x00-\x20]', url)
    except ValueError:
        return False


def parser() -> MarkdownIt:
    md = MarkdownIt('commonmark', {'html': False}).enable('table').enable('strikethrough')
    # Export files must not silently load tracking images or local file URLs.
    md.add_render_rule('image', lambda self, tokens, idx, options, env: escape(
        '[图片：' + tokens[idx].content + '] ' + tokens[idx].attrGet('src')))
    # A structural rule (not raw HTML) keeps citation styling safe for exports.
    def mark_citations(state):
        for block in state.tokens:
            children = block.children or []
            for index, token in enumerate(children):
                if (token.type == 'link_open' and token.attrGet('title') == 'citation'
                        and safe_url(token.attrGet('href') or '')
                        and index + 2 < len(children)
                        and children[index + 1].type == 'text'
                        and re.fullmatch(r'\d+', children[index + 1].content)
                        and children[index + 2].type == 'link_close'):
                    token.meta['citation'] = True
                    children[index + 2].meta['citation'] = True
    md.core.ruler.after('inline', 'citation_links', mark_citations)
    def link_open(self, tokens, idx, options, env):
        return ('<sup class="report-citation">' if tokens[idx].meta.get('citation') else '') + self.renderToken(tokens, idx, options, env)
    def link_close(self, tokens, idx, options, env):
        return self.renderToken(tokens, idx, options, env) + ('</sup>' if tokens[idx].meta.get('citation') else '')
    md.add_render_rule('link_open', link_open)
    md.add_render_rule('link_close', link_close)
    return md


def inline_runs(tokens):
    runs, bold, italic, strike, href, superscript = [], 0, 0, 0, None, False
    for token in tokens or []:
        kind = token.type
        if kind == 'strong_open': bold += 1
        elif kind == 'strong_close': bold -= 1
        elif kind == 'em_open': italic += 1
        elif kind == 'em_close': italic -= 1
        elif kind == 's_open': strike += 1
        elif kind == 's_close': strike -= 1
        elif kind == 'link_open':
            candidate = token.attrGet('href') or ''
            href = candidate if safe_url(candidate) else None
            superscript = bool(token.meta.get('citation'))
        elif kind == 'link_close': href, superscript = None, False
        elif kind in {'text', 'code_inline', 'softbreak', 'hardbreak', 'image', 'html_inline'}:
            text = '\n' if kind in {'softbreak', 'hardbreak'} else token.content
            if kind == 'image': text = f'[图片：{token.content}] {token.attrGet("src") or ""}'
            runs.append({'text': text, 'bold': bool(bold), 'italic': bool(italic),
                         'strike': bool(strike), 'code': kind == 'code_inline', 'href': href,
                         'superscript': superscript})
    return runs


def plain(runs, links=False):
    text, previous_link = '', None
    for run in runs:
        if links and previous_link and previous_link != run.get('href'):
            text += f' ({previous_link})'
        text += run['text']
        previous_link = run.get('href')
    if links and previous_link: text += f' ({previous_link})'
    return text


def parse_blocks(content: str):
    tokens = parser().parse(content)
    blocks, lists, prefix, quote_depth = [], [], '', 0
    i = 0
    while i < len(tokens):
        token = tokens[i]
        if token.type in {'bullet_list_open', 'ordered_list_open'}:
            lists.append({'ordered': token.type == 'ordered_list_open', 'next': int(token.attrGet('start') or 1)})
        elif token.type in {'bullet_list_close', 'ordered_list_close'}:
            lists.pop()
        elif token.type == 'list_item_open':
            current = lists[-1]
            prefix = f"{current['next']}. " if current['ordered'] else '• '
            current['next'] += 1
        elif token.type == 'list_item_close': prefix = ''
        elif token.type == 'blockquote_open': quote_depth += 1
        elif token.type == 'blockquote_close': quote_depth -= 1
        elif token.type == 'table_open':
            rows, row = [], []
            i += 1
            while tokens[i].type != 'table_close':
                cell_token = tokens[i]
                if cell_token.type == 'tr_open': row = []
                elif cell_token.type == 'inline': row.append(inline_runs(cell_token.children))
                elif cell_token.type == 'tr_close': rows.append(row)
                i += 1
            blocks.append({'kind': 'table', 'rows': rows})
        elif token.type == 'inline':
            heading = tokens[i - 1].type == 'heading_open'
            blocks.append({'kind': 'heading' if heading else 'paragraph',
                           'level': int(tokens[i - 1].tag[1:]) if heading else 0,
                           'runs': inline_runs(token.children), 'prefix': prefix,
                           'indent': len(lists), 'quote': quote_depth})
            prefix = ''
        elif token.type in {'fence', 'code_block'}:
            blocks.append({'kind': 'code', 'runs': [{'text': token.content.rstrip('\n')}],
                           'indent': len(lists), 'prefix': prefix})
            prefix = ''
        elif token.type == 'hr': blocks.append({'kind': 'separator'})
        i += 1
    return blocks


def readable_blocks(blocks):
    """Wide, narrative tables become labeled records on portrait paper, losslessly."""
    for block in blocks:
        if block['kind'] != 'table':
            yield block
            continue
        rows = block['rows']
        if len(rows) == 1:
            yield block
            continue
        if len(rows[0]) <= 4 and all(len(plain(cell)) <= 150 for row in rows for cell in row):
            yield block
            continue
        headers = rows[0]
        for index, row in enumerate(rows[1:], 1):
            yield {'kind': 'heading', 'level': 3, 'runs': [{'text': f'表格记录 {index}'}]}
            for header, cell in zip(headers, row):
                yield {'kind': 'paragraph', 'runs': [{'text': plain(header) + '：', 'bold': True}] + cell}


HTML_STYLE = '''
@page { size: A4; margin: 20mm; }
body { margin: 0; color: #252525; background: #f7f3eb; font: 15px/1.85 "Microsoft YaHei", sans-serif; }
article { max-width: 900px; margin: 32px auto; padding: 44px; background: white; }
h1 { text-align: center; font-size: 28px; line-height: 1.5; } h2 { font-size: 21px; margin-top: 28px; } h3 { font-size: 17px; }
p { margin: 0 0 1em; } article:lang(zh) > p { text-indent: 2em; }
li { margin: 7px 0; } a { color: #337556; overflow-wrap: anywhere; }
sup.report-citation { font-size: .72em; line-height: 0; white-space: nowrap; } sup.report-citation a { padding: 0 2px; text-decoration: none; }
table { border-collapse: collapse; width: 100%; margin: 20px 0; font-size: 13px; }
th, td { border: 1px solid #d9d9d9; padding: 10px; vertical-align: top; overflow-wrap: anywhere; }
th { background: #eeeeee; } pre { white-space: pre-wrap; overflow-wrap: anywhere; background: #f3f3f3; padding: 12px; }
blockquote { margin-left: 20px; } h1,h2,h3 { break-after: avoid; } p { orphans: 2; widows: 2; }
@media(max-width: 600px) { article { margin: 0; padding: 20px; overflow-x: auto; } }
@media print { body { background: white; } article { margin: 0; padding: 0; max-width: none; } thead { display: table-header-group; } }
'''


def export_html(topic, content):
    language = 'zh-CN' if re.search(r'[\u3400-\u9fff]', content[:300]) else 'en'
    # HTML is disabled in the parser. CSP additionally blocks active content.
    html = parser().render(content)
    return f'''<!doctype html>
<html lang="{language}"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta http-equiv="Content-Security-Policy" content="default-src 'none'; style-src 'unsafe-inline'; base-uri 'none'; form-action 'none'">
<title>{escape(topic)}</title><style>{HTML_STYLE}</style></head>
<body><article lang="{language}">{html}</article></body></html>'''.encode('utf-8')


def export_docx(topic, blocks):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    doc = Document()
    doc.core_properties.title = topic
    doc.core_properties.author = 'DeepResearch'
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(20)
    section.left_margin = section.right_margin = Mm(22)
    for name, size in [('Normal', 11), ('Title', 20), ('Heading 1', 15), ('Heading 2', 13), ('Heading 3', 12)]:
        style = doc.styles[name]
        style.font.name, style.font.size = 'Times New Roman', Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_after = Pt(8)
    doc.styles['Normal'].paragraph_format.widow_control = True

    def fill(paragraph, runs, prefix=''):
        if prefix: paragraph.add_run(prefix)
        for item in runs:
            run = paragraph.add_run(item['text'])
            run.bold, run.italic, run.font.strike = item.get('bold') or None, item.get('italic') or None, item.get('strike') or None
            run.font.superscript = item.get('superscript') or None
            if item.get('code'): run.font.name = 'Consolas'
            if item.get('href'):
                link = OxmlElement('w:hyperlink')
                link.set(qn('r:id'), paragraph.part.relate_to(item['href'], RT.HYPERLINK, is_external=True))
                run.font.color.rgb = RGBColor.from_string('337556')
                run.underline = True
                link.append(run._r)
                paragraph._p.append(link)

    for block in readable_blocks(blocks):
        kind = block['kind']
        if kind == 'separator':
            doc.add_paragraph()
        elif kind == 'table':
            rows = block['rows']
            table = doc.add_table(rows=1, cols=len(rows[0]))
            table.style = 'Table Grid'
            table.autofit = False
            weights = [max(6, min(35, max(len(plain(row[col])) for row in rows))) for col in range(len(rows[0]))]
            for col, weight in zip(table.columns, weights): col.width = Mm(166 * weight / sum(weights))
            for index, row in enumerate(rows):
                cells = table.rows[0].cells if index == 0 else table.add_row().cells
                for column, (cell, runs) in enumerate(zip(cells, row)):
                    cell.width = Mm(166 * weights[column] / sum(weights))
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    fill(cell.paragraphs[0], runs)
                    cell.paragraphs[0].paragraph_format.space_after = Pt(5)
                    cell.paragraphs[0].paragraph_format.space_before = Pt(5)
                    tcpr = cell._tc.get_or_add_tcPr()
                    borders = OxmlElement('w:tcBorders')
                    for side in ('top', 'left', 'bottom', 'right'):
                        border = OxmlElement('w:' + side)
                        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]: border.set(qn('w:' + key), value)
                        borders.append(border)
                    tcpr.append(borders)
                    if index == 0:
                        shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), 'EEEEEE'); tcpr.append(shading)
                        for run in cell.paragraphs[0].runs: run.bold = True
                if index == 0:
                    repeat = OxmlElement('w:tblHeader')
                    table.rows[0]._tr.get_or_add_trPr().append(repeat)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        else:
            level = block.get('level', 0)
            style = ('Title' if level == 1 else f'Heading {min(level - 1, 3)}') if kind == 'heading' else 'Normal'
            paragraph = doc.add_paragraph(style=style)
            fill(paragraph, block['runs'], block.get('prefix', ''))
            if kind == 'heading' and level == 1: paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif block.get('indent') or block.get('quote'):
                paragraph.paragraph_format.left_indent = Mm(5 * max(block.get('indent', 0), block.get('quote', 0)))
            elif kind == 'paragraph' and re.search(r'[\u3400-\u9fff]', plain(block['runs'])):
                paragraph.paragraph_format.first_line_indent = Pt(22)
            if kind == 'code':
                paragraph.paragraph_format.line_spacing = 1.1
                for run in paragraph.runs: run.font.name, run.font.size = 'Consolas', Pt(9)
    stream = BytesIO(); doc.save(stream)
    return stream.getvalue()


@lru_cache(maxsize=4)
def pdf_font(font_path):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    candidate = font_path or ('C:/Windows/Fonts/simsun.ttc' if os.name == 'nt' else '')
    if candidate:
        if not Path(candidate).is_file(): raise ValueError('PDF 中文字体文件不存在，请检查 EXPORT_FONT_PATH')
        name = 'ReportCJK' + str(abs(hash(candidate)))
        pdfmetrics.registerFont(TTFont(name, candidate, subfontIndex=0))
    else:
        name = 'STSong-Light'
        pdfmetrics.registerFont(UnicodeCIDFont(name))
    pdfmetrics.registerFontFamily(name, normal=name, bold=name, italic=name, boldItalic=name)
    return name


def export_pdf(topic, blocks, font_path=''):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle

    font = pdf_font(font_path)
    body = ParagraphStyle('Body', fontName=font, fontSize=10.5, leading=17, spaceAfter=9, wordWrap='CJK')
    styles = {level: ParagraphStyle(f'H{level}', parent=body, fontSize=size, leading=size * 1.5,
                                   spaceBefore=14, spaceAfter=10, keepWithNext=True,
                                   alignment=TA_CENTER if level == 1 else 0)
              for level, size in [(1, 20), (2, 15), (3, 12), (4, 11), (5, 11), (6, 11)]}
    cell_style = ParagraphStyle('Cell', parent=body, fontSize=9, leading=14, spaceAfter=0)
    stream, story = BytesIO(), []
    document = SimpleDocTemplate(stream, pagesize=A4, rightMargin=22*mm, leftMargin=22*mm,
                                 topMargin=20*mm, bottomMargin=20*mm, title=topic, author='DeepResearch')

    def markup(runs):
        parts = []
        for run in runs:
            text = escape(run['text']).replace('\n', '<br/>')
            if run.get('bold'): text = f'<b>{text}</b>'
            if run.get('italic'): text = f'<i>{text}</i>'
            if run.get('strike'): text = f'<strike>{text}</strike>'
            if run.get('href'): text = f'<link href="{escape(run["href"], quote=True)}" color="#337556">{text}</link>'
            if run.get('superscript'): text = f'<super>{text}</super>'
            parts.append(text)
        return ''.join(parts)

    for block in readable_blocks(blocks):
        kind = block['kind']
        if kind == 'separator': story.append(Spacer(1, 10))
        elif kind == 'table':
            rows = block['rows']
            weights = [max(6, min(35, max(len(plain(row[col])) for row in rows))) for col in range(len(rows[0]))]
            table = LongTable([[Paragraph(markup(cell), cell_style) for cell in row] for row in rows],
                              colWidths=[document.width * weight / sum(weights) for weight in weights],
                              repeatRows=1, splitByRow=1, splitInRow=1)
            table.setStyle(TableStyle([('GRID', (0,0), (-1,-1), .5, colors.HexColor('#D9D9D9')),
                                       ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#EEEEEE')),
                                       ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
                                       ('TOPPADDING', (0,0), (-1,-1), 7), ('BOTTOMPADDING', (0,0), (-1,-1), 7)]))
            story.extend([table, Spacer(1, 12)])
        else:
            style = styles[block['level']] if kind == 'heading' else body
            if kind != 'heading':
                indent = max(block.get('indent', 0), block.get('quote', 0)) * 12
                style = ParagraphStyle('Paragraph', parent=body, leftIndent=indent,
                                       firstLineIndent=21 if kind == 'paragraph' and not indent and re.search(r'[\u3400-\u9fff]', plain(block['runs'])) else 0)
            text = escape(block.get('prefix', '')) + markup(block['runs'])
            if kind == 'code':
                text = text.replace(' ', '&#160;')
                style = ParagraphStyle('Code', parent=body, fontSize=9, leading=13, leftIndent=8)
            story.append(Paragraph(text or ' ', style))
    if not story: story.append(Paragraph(escape(topic), body))
    document.build(story)
    return stream.getvalue()


def export_report(topic, content, format, font_path=''):
    if format not in MIME_TYPES: raise ValueError('不支持的导出格式')
    if format == 'md': return content.encode('utf-8')
    if format == 'html': return export_html(topic, content)
    blocks = parse_blocks(content)
    if format == 'docx': return export_docx(topic, blocks)
    if format == 'pdf': return export_pdf(topic, blocks, font_path)
    parts = []
    for block in blocks:
        if block['kind'] == 'table':
            parts.append('\n'.join('\t'.join(plain(cell, links=True) for cell in row) for row in block['rows']))
        elif block['kind'] != 'separator': parts.append(block.get('prefix', '') + plain(block['runs'], links=True))
    return '\n\n'.join(parts).encode('utf-8-sig')
